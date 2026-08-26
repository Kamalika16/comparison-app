
from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SEVERITY_FILL = {
    "HIGH": PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid"),
    "MEDIUM": PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid"),
    "LOW": PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid"),
}

HEADER_FILL = PatternFill(start_color="305496", end_color="305496", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def detail_headers(meta: dict | None) -> list[str]:
    """Column layout mirrors the on-screen results table (Phase 6). The first
    header embeds the dynamically selected primary identifier column."""
    att_key = (meta or {}).get("attendance_key_column")
    id_header = (
        f"Primary Identifier ({att_key})" if att_key else "Primary Identifier"
    )
    return [
        id_header,
        "Employee Name",
        "Attendance Hours",
        "Client Hours",
        "Difference",
        "Classification",
        "Severity",
        "Reason",
        "Recommendation",
    ]


def _difference(m: dict):
    """|Attendance - Client| rounded to 2 dp — identical to the value the UI
    shows; blank when either side has no value."""
    a, b = m.get("company_hours"), m.get("client_hours")
    if a is None or b is None:
        return None
    try:
        return round(abs(float(a) - float(b)), 2)
    except (TypeError, ValueError):
        return None


def _classification_label(m: dict) -> str:
    """Human-readable classification identical to the UI badge, derived from
    the same null checks (never invented)."""
    has_att = m.get("company_hours") is not None
    has_cli = m.get("client_hours") is not None
    if has_att and not has_cli:
        return "Only in iLink Attendance"
    if has_cli and not has_att:
        return "Only in Client Worksheet"
    if has_att and has_cli:
        return "Hours mismatch"
    return "Incomplete record"


def _autosize_columns(ws, headers) -> None:
    for idx, header in enumerate(headers, start=1):
        col_letter = get_column_letter(idx)
        max_len = max(
            [len(str(header))]
            + [len(str(cell.value)) for cell in ws[col_letter] if cell.value is not None]
        )
        ws.column_dimensions[col_letter].width = min(max(max_len + 2, 12), 45)


def write_report(result: dict, output_path: str, meta: dict | None = None) -> Path:
    summary = result.get("summary", {})
    mismatches = result.get("mismatches", [])
    meta = meta or {}
    headers = detail_headers(meta)

    wb = Workbook()

    # --- Summary sheet (metrics + factual export metadata) ---
    ws_summary = wb.active
    ws_summary.title = "Summary"
    ws_summary.append(["Metric", "Value"])
    for cell in ws_summary[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT

    summary_rows = [
        ("Generated (UTC)", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M")),
        ("Total Records Compared", summary.get("total_records_compared", 0)),
        ("Matches", summary.get("matches", 0)),
        ("Mismatches", summary.get("mismatches", 0)),
        ("Needs Review (High Severity)", summary.get("review_required", 0)),
        ("Attendance Primary Identifier", meta.get("attendance_key_column", "n/a")),
        ("Client Primary Identifier", meta.get("client_key_column", "n/a")),
        ("Attendance Hours Column", meta.get("attendance_hours_column", "n/a")),
        ("Client Hours Column", meta.get("client_hours_column", "n/a")),
    ]
    for label, value in summary_rows:
        ws_summary.append([label, value])

    ws_summary.column_dimensions["A"].width = 32
    ws_summary.column_dimensions["B"].width = 34

    # --- Details sheet (mismatches only — mirrors the UI table) ---
    ws_details = wb.create_sheet("Details")
    ws_details.append(headers)
    for cell in ws_details[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")

    for m in mismatches:
        ws_details.append([
            m.get("employee_id", "") or "",
            m.get("employee_name", "") or "",
            m.get("company_hours"),
            m.get("client_hours"),
            _difference(m),
            _classification_label(m),
            m.get("severity", ""),
            m.get("reason", "") or "",
            m.get("recommendation", "") or "",
        ])
        fill = SEVERITY_FILL.get(m.get("severity"))
        if fill:
            for cell in ws_details[ws_details.max_row]:
                cell.fill = fill

    ws_details.freeze_panes = "A2"
    ws_details.auto_filter.ref = ws_details.dimensions
    _autosize_columns(ws_details, headers)

    output_path = Path(output_path)
    wb.save(output_path)
    return output_path